"""
Export engine — CSV, XLSX, JSON, XML, PDF.
Usage: call export_queryset(request, qs, fields, filename, fmt)
"""
import csv
import io
import json
from xml.etree import ElementTree as ET

from django.http import HttpResponse


def _coerce(val):
    if val is None:
        return ""
    if isinstance(val, (list, dict)):
        return json.dumps(val)
    return str(val)


def export_queryset(request, rows: list, fields: list, filename: str, fmt: str) -> HttpResponse:
    """
    rows   — list of dicts (each dict is one row)
    fields — list of (header_label, key) tuples
    fmt    — 'csv' | 'xlsx' | 'json' | 'xml' | 'pdf'
    """
    fmt = fmt.lower()
    if fmt == "csv":
        return _export_csv(rows, fields, filename)
    if fmt == "xlsx":
        return _export_xlsx(rows, fields, filename)
    if fmt == "json":
        return _export_json(rows, fields, filename)
    if fmt == "xml":
        return _export_xml(rows, fields, filename)
    if fmt == "pdf":
        return _export_pdf(rows, fields, filename)
    return HttpResponse("Unknown format", status=400)


def _export_csv(rows, fields, filename):
    resp = HttpResponse(content_type="text/csv")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.csv"'
    w = csv.writer(resp)
    w.writerow([h for h, _ in fields])
    for row in rows:
        w.writerow([_coerce(row.get(k, "")) for _, k in fields])
    return resp


def _export_xlsx(rows, fields, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    ws = wb.active
    ws.title = filename[:31]

    header_font = Font(bold=True, color="00FF41", name="Courier New")
    header_fill = PatternFill("solid", fgColor="0D1F0D")

    for col, (header, _) in enumerate(fields, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left")

    for r, row in enumerate(rows, 2):
        for col, (_, key) in enumerate(fields, 1):
            ws.cell(row=r, column=col, value=_coerce(row.get(key, "")))

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    resp = HttpResponse(
        buf.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}.xlsx"'
    return resp


def _export_json(rows, fields, filename):
    data = [{h: _coerce(row.get(k, "")) for h, k in fields} for row in rows]
    resp = HttpResponse(json.dumps(data, indent=2), content_type="application/json")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.json"'
    return resp


def _export_xml(rows, fields, filename):
    root = ET.Element("records")
    for row in rows:
        rec = ET.SubElement(root, "record")
        for h, k in fields:
            el = ET.SubElement(rec, h.lower().replace(" ", "_"))
            el.text = _coerce(row.get(k, ""))
    buf = ET.tostring(root, encoding="unicode", xml_declaration=False)
    resp = HttpResponse(f'<?xml version="1.0" encoding="UTF-8"?>\n{buf}', content_type="application/xml")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.xml"'
    return resp


def _export_pdf(rows, fields, filename):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=20, rightMargin=20, topMargin=20, bottomMargin=20)

    styles = getSampleStyleSheet()
    black = colors.HexColor("#000000")
    green = colors.HexColor("#00ff41")
    dkgreen = colors.HexColor("#003a10")
    ltgreen = colors.HexColor("#ccffcc")

    table_data = [[h for h, _ in fields]]
    for row in rows:
        table_data.append([_coerce(row.get(k, "")) for _, k in fields])

    col_count = len(fields)
    col_width = (landscape(A4)[0] - 40) / col_count

    t = Table(table_data, colWidths=[col_width] * col_count, repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND",  (0, 0), (-1, 0),  dkgreen),
        ("TEXTCOLOR",   (0, 0), (-1, 0),  green),
        ("FONTNAME",    (0, 0), (-1, 0),  "Courier-Bold"),
        ("FONTSIZE",    (0, 0), (-1, 0),  8),
        ("BACKGROUND",  (0, 1), (-1, -1), black),
        ("TEXTCOLOR",   (0, 1), (-1, -1), ltgreen),
        ("FONTNAME",    (0, 1), (-1, -1), "Courier"),
        ("FONTSIZE",    (0, 1), (-1, -1), 7),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [black, colors.HexColor("#080808")]),
        ("GRID",        (0, 0), (-1, -1), 0.3, colors.HexColor("#003a10")),
        ("VALIGN",      (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING",  (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))

    title = Paragraph(f'<font color="#00ff41" name="Courier-Bold">{filename}</font>', styles["Normal"])
    doc.build([title, t])
    buf.seek(0)
    resp = HttpResponse(buf.read(), content_type="application/pdf")
    resp["Content-Disposition"] = f'attachment; filename="{filename}.pdf"'
    return resp
